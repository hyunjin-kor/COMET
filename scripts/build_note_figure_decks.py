"""Author the 2026-09-16 figure decks from artwork, panel renders and native labels.

This is the initial authoring tool. The tracked PowerPoint files in
docs/paper/diagram-sources-2026-09-16 are the editable sources afterwards; do not
rerun this script over edited decks unless the numeric panels changed. Panels come
from ``draw_application_note_figures.py --panels``; the Figure 2(a) and 3(a) labels
are transplanted from the 2026-09-13-h26 and 2026-09-15 decks. Requires python-pptx.
After building, export with scripts/export_note_diagram_slides.ps1.
"""

import argparse
import json
import sys
from copy import deepcopy
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Mm, Pt

ROOT = Path(__file__).resolve().parents[1]
DECKS = ROOT / "docs/paper/diagram-sources-2026-09-16"
PANELS = DECKS / "panels"
ARTWORK = DECKS / "artwork"
H26 = ROOT / "docs/paper/diagram-sources-2026-09-13-h26"
H45 = ROOT / "docs/paper/diagram-sources-2026-09-15"
INK = RGBColor(0x1F, 0x2A, 0x30)
FONT = {"en": "Arial", "ko": "Malgun Gothic"}
LANGS = ("en", "ko")
NOTE = ("Editable COMET figure deck, 2026-09-16. Numeric panels are matplotlib renders of the frozen "
        "COMET analysis files (see panels/panels.json); they are not editable data. Conceptual artwork is "
        "identified in artwork-provenance.json. Labels and panel letters are native slide objects.")
LABELS = {
    "figS1_allocation": {
        "en": ["Solution aliquot\n(volume fraction)", "Intermediate transfer\n(mass fraction)", "Final batch\n(dry mass)"],
        "ko": ["용액 분취\n(부피 분율)", "중간체 이전\n(질량 분율)", "최종 배치\n(건조 질량)"],
        "centers": (30, 89, 153),
    },
    "figS8_provenance": {
        "en": ["Source\n(DOI and locator)", "Record\n(reported value, user edit)", "Cost contribution\n(checksum)"],
        "ko": ["출처\n(DOI·원문 위치)", "기록\n(보고값·수정값)", "비용 기여분\n(체크섬)"],
        "centers": (37, 93, 150),
    },
    "figS9_interface": {
        "en": ["Preparation record with source-linked inputs", "Batch cost contributions by operation"],
        "ko": ["출처가 연결된 입력을 담은 제조 기록", "조작별 배치 비용 기여분"],
    },
}


def load_panels():
    manifest = json.loads((PANELS / "panels.json").read_text(encoding="utf-8"))
    records = {}
    for row in manifest["panels"]:
        records.setdefault((row["figure"], row["language"]), {})[row["panel"]] = row
    return records


def new_deck(width_mm, height_mm):
    deck = Presentation()
    deck.slide_width, deck.slide_height = Mm(width_mm), Mm(height_mm)
    return deck


def blank_slide(deck):
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    # A full-slide white rectangle keeps PowerPoint's SVG viewBox equal to the slide.
    background = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, deck.slide_width, deck.slide_height)
    background.name = "slide-background"
    background.fill.solid()
    background.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    background.line.fill.background()
    background.shadow.inherit = False
    slide.notes_slide.notes_text_frame.text = NOTE
    return slide


def picture(slide, path, x, y, width, height, name):
    shape = slide.shapes.add_picture(str(path), Mm(x), Mm(y), Mm(width), Mm(height))
    shape.name = name
    return shape


def label(slide, text, x, y, width, height, size, lang, bold=False, align=PP_ALIGN.CENTER, name="label"):
    box = slide.shapes.add_textbox(Mm(x), Mm(y), Mm(width), Mm(height))
    box.name = name
    frame = box.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    for index, line in enumerate(text.split("\n")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = align
        run = paragraph.add_run()
        run.text = line
        run.font.name = "Arial" if line.isascii() else FONT[lang]
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = INK
    return box


def letter(slide, char, x, y):
    label(slide, f"({char})", x, y, 7, 5, 10, "en", bold=True, align=PP_ALIGN.LEFT, name=f"panel-{char}")


def next_shape_id(slide):
    ids = [int(node.get("id")) for node in slide.shapes._spTree.xpath(".//p:cNvPr")]
    return max(ids, default=1) + 1


def transplant(source, slide_index, target, dx, dy, scale):
    """Copy artwork and text objects from an earlier deck slide, scaled and offset in mm."""
    source_slide = Presentation(source).slides[slide_index]
    for shape in source_slide.shapes:
        text = shape.text_frame.text.strip() if shape.has_text_frame else ""
        if shape.shape_type == 13:
            new = target.shapes.add_picture(_blob(shape), Mm(shape.left / 36000 * scale + dx), Mm(shape.top / 36000 * scale + dy),
                                            int(shape.width * scale), int(shape.height * scale))
            new.name = shape.name or "artwork"
            new.crop_left, new.crop_top = shape.crop_left, shape.crop_top
            new.crop_right, new.crop_bottom = shape.crop_right, shape.crop_bottom
        elif text and text != "(a)":
            element = deepcopy(shape._element)
            for node in element.xpath(".//p:cNvPr"):
                node.set("id", str(next_shape_id(target)))
            target.shapes._spTree.append(element)
            new = target.shapes[-1]
            new.left, new.top = Mm(shape.left / 36000 * scale + dx), Mm(shape.top / 36000 * scale + dy)
            new.width, new.height = int(shape.width * scale), int(shape.height * scale)
            for paragraph in new.text_frame.paragraphs:
                for run in paragraph.runs:
                    if run.font.size is not None:
                        run.font.size = Pt(round(run.font.size.pt * scale * 2) / 2)


def _blob(shape):
    import io

    return io.BytesIO(shape.image.blob)


def place(slide, records, panel, x, y):
    row = records[panel]
    picture(slide, PANELS / row["file"], x, y, row["width_mm"], row["height_mm"], f"panel-{panel}-image")


def build_fig2(records):
    deck = new_deck(178, 212)
    for index, lang in enumerate(LANGS):
        slide = blank_slide(deck)
        transplant(H26 / "fig2a_cost_model.pptx", index, slide, 0, 2, 1.0)
        place(slide, records[("fig2_cost_model", lang)], "b", 0, 64)
        place(slide, records[("fig2_cost_model", lang)], "c", 0, 166)
        letter(slide, "a", 2, 0.5)
        letter(slide, "b", 2, 64.5)
        letter(slide, "c", 2, 166.5)
    return deck


def build_fig3(records):
    deck = new_deck(178, 142)
    for index, lang in enumerate(LANGS):
        slide = blank_slide(deck)
        transplant(H45 / "fig3a_manufacturing_v2.pptx", index, slide, 0, 1.7, 178 / 370.417)
        place(slide, records[("fig3_manufacturing", lang)], "b", 0, 62)
        place(slide, records[("fig3_manufacturing", lang)], "c", 89, 62)
        letter(slide, "a", 2, 0.5)
        letter(slide, "b", 2, 64)
        letter(slide, "c", 91, 64)
    return deck


def build_fig4(records):
    deck = new_deck(178, 128)
    for lang in LANGS:
        slide = blank_slide(deck)
        rows = records[("fig4_ranking", lang)]
        for panel, x, y in (("a", 0, 0), ("b", 59, 0), ("c", 118, 0), ("d", 0, 64), ("e", 89, 64)):
            place(slide, rows, panel, x, y)
            letter(slide, panel, x + 2, y + 0.5)
    return deck


def build_ranking(records):
    deck = new_deck(178, 203)
    for lang in LANGS:
        slide = blank_slide(deck)
        rows = records[("figS6_ranking_tests", lang)]
        for panel, x, y in (("a", 0, 0), ("b", 0, 132), ("c", 89, 132)):
            place(slide, rows, panel, x, y)
            letter(slide, panel, x + 2, y + 0.5)
    return deck


def build_s1(records):
    deck = new_deck(84, 122)
    for lang in LANGS:
        slide = blank_slide(deck)
        rows = records[("figS4_metal_prices", lang)]
        place(slide, rows, "a", 0, 0)
        place(slide, rows, "b", 0, 60)
        letter(slide, "a", 1, 0.5)
        letter(slide, "b", 1, 60.5)
    return deck


def build_single(name, width, height, records):
    deck = new_deck(width, height)
    for lang in LANGS:
        slide = blank_slide(deck)
        place(slide, records[(name, lang)], "a", 0, 0)
    return deck


def build_s3(records):
    deck = new_deck(178, 72)
    for lang in LANGS:
        slide = blank_slide(deck)
        rows = records[("figS3_monte_carlo", lang)]
        place(slide, rows, "a", 0, 0)
        place(slide, rows, "b", 89, 0)
        letter(slide, "a", 2, 0.5)
        letter(slide, "b", 91, 0.5)
    return deck


def build_s5(records):
    deck = new_deck(178, 88.5)
    for lang in LANGS:
        slide = blank_slide(deck)
        rows = records[("figS5_crossovers", lang)]
        place(slide, rows, "a", 0, 0)
        place(slide, rows, "b", 87, 0)
        letter(slide, "a", 2, 0.5)
        letter(slide, "b", 89, 0.5)
    return deck


def build_concept(name):
    deck = new_deck(178, 70)
    spec = LABELS[name]
    for lang in LANGS:
        slide = blank_slide(deck)
        picture(slide, ARTWORK / f"{name}.text-free.png", 0, 0, 178, 178 * 592 / 1792, "artwork")
        for center, text in zip(spec["centers"], spec[lang], strict=True):
            label(slide, text, center - 28, 59.5, 56, 10, 9.5, lang, name=f"stage-{spec[lang].index(text)}")
    return deck


def build_s6():
    """Two interface views side by side with native captions; screenshots are listed in figS9_interface.capture.json."""
    from PIL import Image

    views = ["figS9_interface.record", "figS9_interface.costs"]
    width = 87
    heights = []
    for view in views:
        with Image.open(ARTWORK / f"{view}.en.png") as image:
            heights.append(width * image.height / image.width)
    deck = new_deck(178, max(heights) + 12)
    spec = LABELS["figS9_interface"]
    for lang in LANGS:
        slide = blank_slide(deck)
        for index, (view, height) in enumerate(zip(views, heights, strict=True)):
            x = index * (width + 4)
            picture(slide, ARTWORK / f"{view}.{lang}.png", x, 11, width, height, f"view-{index}")
            letter(slide, "ab"[index], x + 1, 0.5)
            label(slide, spec[lang][index], x + 9, 0.5, width - 9, 9, 8.5, lang, align=PP_ALIGN.LEFT, name=f"caption-{index}")
    return deck


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("names", nargs="*", help="Deck names to build (default: all)")
    args = parser.parse_args()
    records = load_panels()
    builders = {
        "fig2_cost_model": lambda: build_fig2(records),
        "fig3_manufacturing": lambda: build_fig3(records),
        "fig4_ranking": lambda: build_fig4(records),
        "figS4_metal_prices": lambda: build_s1(records),
        "figS2_sensitivity": lambda: build_single("figS2_sensitivity", 150, 82, records),
        "figS3_monte_carlo": lambda: build_s3(records),
        "figS7_evidence": lambda: build_single("figS7_evidence", 178, 118, records),
        "figS5_crossovers": lambda: build_s5(records),
        "figS6_ranking_tests": lambda: build_ranking(records),
        "figS9_interface": build_s6,
        "figS1_allocation": lambda: build_concept("figS1_allocation"),
        "figS8_provenance": lambda: build_concept("figS8_provenance"),
    }
    names = args.names or list(builders)
    unknown = [name for name in names if name not in builders]
    if unknown:
        sys.exit(f"unknown deck(s): {unknown}")
    for name in names:
        deck = builders[name]()
        destination = DECKS / f"{name}.pptx"
        deck.save(destination)
        print("wrote", destination)


if __name__ == "__main__":
    main()
